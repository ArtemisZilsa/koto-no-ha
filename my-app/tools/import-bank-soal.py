"""Excel bank soal (Drive Zilsa) -> supabase/migrations/061_seed_bank_soal.sql.

Pakai:  python tools/import-bank-soal.py <folder-berisi-xlsx>
Butuh:  pip install openpyxl

Membaca setiap file Bank-Soal-*.xlsx yang punya sheet 'Bank Soal' (tata bahasa)
atau 'Bank Soal - Kanji'. Upsert per kode soal, jadi aman diulang setelah Excel
direvisi. Baris berstatus 'Buang' dihapus dari database.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent.parent / 'supabase/migrations/061_seed_bank_soal.sql'
KEYS = 'ABCD'
MODES = {'Latihan': 'latihan', 'Checkpoint': 'checkpoint', 'Cadangan': 'cadangan'}


def status(raw):
    s = (raw or '').strip().lower()
    if s == 'buang':
        return None
    if s == 'ok':
        return 'ok'
    if s == 'revisi':
        return 'revisi'
    return 'belum'  # 'Belum', 'Belum direview', kosong


def rows(ws):
    head = [str(h).strip() if h else '' for h in next(ws.iter_rows(max_row=1, values_only=True))]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0] is not None:
            yield dict(zip(head, r))


def soal(r, **kw):
    opts = [str(r[f'Opsi {k}']).strip() for k in KEYS]
    key = str(r['Kunci']).strip().upper()
    assert all(opts), f"{kw['code']}: opsi kosong"
    assert len(set(opts)) == 4, f"{kw['code']}: opsi kembar"
    assert key in KEYS, f"{kw['code']}: kunci '{key}' bukan A-D"
    return dict(kw, question=str(r['Soal']).strip(), options=opts, answer_index=KEYS.index(key),
                explanation=r.get('Penjelasan'), distractor_basis=r.get('Basis Distraktor'))


def read(folder):
    out, drop = [], []
    for f in sorted(Path(folder).glob('Bank-Soal-*.xlsx')):
        wb = load_workbook(f, data_only=True)
        if 'Bank Soal' in wb.sheetnames:  # tata bahasa
            for i, r in enumerate(rows(wb['Bank Soal'])):
                code, st = str(r['ID Soal']).strip(), status(r['Status Review'])
                if st is None:
                    drop.append(code)
                    continue
                out.append(soal(r, code=code, level=r['Level'], category='tata_bahasa', qtype=r['Tipe'],
                                unit=int(r['Batch']), group_code=r['ID Pola'], group_label=r['Pola'],
                                mode=MODES[r['Mode']], checkpoint=r['Checkpoint'], review_status=st, order_index=i))
        if 'Bank Soal - Kanji' in wb.sheetnames:
            level = 'N5' if '-N5-' in f.name else sys.exit(f'{f.name}: level tidak dikenali dari nama file')
            for r in rows(wb['Bank Soal - Kanji']):
                code, st = f"K{level[1]}-{int(r['No']):04d}", status(r['Status'])
                if st is None:
                    drop.append(code)
                    continue
                t = int(r['Level'])
                out.append(soal(r, code=code, level=level, category='kanji', qtype=r['Tipe'], unit=t,
                                group_code=f'T{t:02d}', group_label=r['Tema Level'], mode='latihan',
                                checkpoint=None, review_status=st, order_index=int(r['No'])))
    return out, drop


def lit(v):
    if v is None:
        return 'NULL'
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, list):
        return lit(json.dumps(v, ensure_ascii=False)) + '::jsonb'
    return "'" + str(v).strip().replace("'", "''") + "'"


COLS = ['code', 'level', 'category', 'qtype', 'unit', 'group_code', 'group_label', 'mode', 'checkpoint',
        'question', 'options', 'answer_index', 'explanation', 'distractor_basis', 'review_status', 'order_index']


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    data, drop = read(sys.argv[1])
    codes = [d['code'] for d in data]
    assert len(codes) == len(set(codes)), 'kode soal kembar'
    sql = ['-- 061: Seed bank soal. DIBUAT OTOMATIS oleh tools/import-bank-soal.py, jangan diedit manual.',
           f'-- {len(data)} soal. Upsert per code: aman dijalankan ulang.', 'BEGIN;']
    for i in range(0, len(data), 200):
        vals = ',\n'.join('(' + ', '.join(lit(d[c]) for c in COLS) + ')' for d in data[i:i + 200])
        upd = ', '.join(f'{c} = EXCLUDED.{c}' for c in COLS[1:])
        sql.append(f'INSERT INTO public.bank_soal ({", ".join(COLS)}) VALUES\n{vals}\nON CONFLICT (code) DO UPDATE SET {upd};')
    if drop:
        sql.append(f'DELETE FROM public.bank_soal WHERE code IN ({", ".join(lit(c) for c in drop)});')
    sql.append('COMMIT;')
    OUT.write_text('\n'.join(sql) + '\n', encoding='utf-8')
    by = {}
    for d in data:
        by[(d['category'], d['mode'])] = by.get((d['category'], d['mode']), 0) + 1
    print(f'{len(data)} soal -> {OUT.name}; dibuang {len(drop)}; {by}')


if __name__ == '__main__':
    main()
